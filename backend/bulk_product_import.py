import io
import re
import uuid
from typing import Any

from fastapi import HTTPException

from db import db
from rates import get_rate
from services import now_iso, notify_product_created
from inventory import add_records, available_count


PRODUCT_SHEET_NAME = "product"
SERVICE_WAIT_VALUES = {1, 5, 10, 25, 60}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value)).strip().casefold()


def _parse_num(value: Any, idr: bool = False) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    s = _text(value).replace("Rp", "").replace("$", "").strip()
    if idr:
        s = s.replace(".", "").replace(",", "")
    else:
        s = s.replace(",", "")
    return float(s)


def _generate_description(name: str, points: str, kind: str, wait_minutes: int | None = None) -> str:
    name = name.strip()
    points = points.strip()

    intro = points or (
        f"{name} merupakan produk digital yang siap digunakan untuk kebutuhan digital Anda."
        if kind == "digital"
        else f"{name} merupakan layanan yang diproses oleh tim kami setelah pesanan berhasil dibuat."
    )

    if kind == "digital":
        return (
            f"{name}\n\n"
            f"{intro}\n\n"
            "Keunggulan:\n"
            "• Detail produk diproses secara otomatis setelah pembayaran berhasil.\n"
            "• Pengiriman dilakukan melalui bot sehingga proses lebih cepat dan praktis.\n"
            "• Data produk mengikuti inventory yang tersedia pada sistem.\n\n"
            "Catatan:\n"
            "• Pastikan detail pesanan sudah sesuai sebelum melakukan pembayaran.\n"
            "• Jika terdapat kendala pada produk yang diterima, silakan hubungi admin."
        )

    wait = wait_minutes or 5
    return (
        f"{name}\n\n"
        f"{intro}\n\n"
        "Informasi layanan:\n"
        f"• Pesanan diproses setelah pembayaran berhasil.\n"
        f"• Estimasi waktu tunggu: {wait} menit.\n"
        "• Status dan informasi proses akan diberikan melalui bot.\n\n"
        "Catatan:\n"
        "• Waktu proses dapat bergantung pada antrean layanan.\n"
        "• Jika terdapat kendala, silakan hubungi admin."
    )


def _column_map(headers: list[str]) -> dict[str, int]:
    aliases = {
        "name": ["nama product", "nama produk", "product", "produk", "name"],
        "short_description": [
            "deskripsi singkat", "short description", "poin deskripsi",
            "point deskripsi", "deskripsi", "description", "desc",
        ],
        "price_usd": ["harga usd", "price usd", "usd", "price"],
        "price_idr": ["harga idr", "price idr", "idr"],
        "kind": ["jenis product", "jenis produk", "product kind", "tipe product", "tipe produk"],
        "wait": ["waktu tunggu (menit)", "waktu tunggu", "wait minutes", "service wait minutes"],
        "message": ["pesan jasa", "pesan antrean jasa", "service message", "service message template"],
    }
    normalized = {_norm(h): i for i, h in enumerate(headers)}
    result = {}
    for key, names in aliases.items():
        for alias in names:
            if _norm(alias) in normalized:
                result[key] = normalized[_norm(alias)]
                break
    return result


def _parse_product_rows(ws, rate: float):
    rows = list(ws.iter_rows(values_only=True))
    rows = [list(row) for row in rows if any(_text(v) for v in row)]
    if not rows:
        raise HTTPException(400, 'Sheet "product" kosong.')

    headers = [_text(v) for v in rows[0]]
    cmap = _column_map(headers)

    for required in ("name", "kind"):
        if required not in cmap:
            raise HTTPException(400, f'Sheet "product" wajib memiliki kolom: Nama Product dan Jenis Product.')

    if "price_usd" not in cmap and "price_idr" not in cmap:
        raise HTTPException(400, 'Sheet "product" wajib memiliki kolom Harga USD atau Harga IDR.')

    products = []
    errors = []

    for row_no, row in enumerate(rows[1:], start=2):
        def cell(key):
            idx = cmap.get(key)
            return _text(row[idx]) if idx is not None and idx < len(row) else ""

        name = cell("name")
        if not name:
            continue

        kind_raw = cell("kind").casefold()
        if kind_raw.startswith("a.") or "digital" in kind_raw or "data" in kind_raw:
            kind = "digital"
        elif kind_raw.startswith("b.") or "jasa" in kind_raw or "service" in kind_raw:
            kind = "service"
        else:
            errors.append(f"Baris {row_no}: Jenis Product tidak valid.")
            continue

        try:
            usd_raw = cell("price_usd")
            idr_raw = cell("price_idr")
            price_usd = _parse_num(usd_raw, False) if usd_raw else 0.0
            price_idr = _parse_num(idr_raw, True) if idr_raw else None
            if price_usd <= 0 and (price_idr is None or price_idr <= 0):
                raise ValueError("harga kosong")
            if price_usd <= 0:
                price_usd = round(price_idr / rate, 2)
            if price_idr is not None and price_idr <= 0:
                price_idr = None
        except (ValueError, TypeError):
            errors.append(f"Baris {row_no}: harga tidak valid.")
            continue

        wait_minutes = None
        service_message = ""
        if kind == "service":
            wait_raw = cell("wait")
            try:
                wait_minutes = int(float(wait_raw)) if wait_raw else 5
            except ValueError:
                wait_minutes = 5
            if wait_minutes not in SERVICE_WAIT_VALUES:
                errors.append(f"Baris {row_no}: Waktu Tunggu harus 1, 5, 10, 25, atau 60 menit.")
                continue
            service_message = cell("message") or (
                "Jasa {product_name} sedang dalam antrean, harap tunggu "
                "{wait_minutes} menit untuk dapat menghubungi admin."
            )

        points = cell("short_description")
        description = _generate_description(name, points, kind, wait_minutes)

        products.append({
            "row_no": row_no,
            "name": name,
            "description": description,
            "short_description": points,
            "price_usd": round(price_usd, 2),
            "price_idr": price_idr,
            "product_kind": kind,
            "wait_minutes": wait_minutes,
            "service_message": service_message,
        })

    return products, errors


def _inventory_from_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    rows = [list(row) for row in rows if any(_text(v) for v in row)]
    if not rows:
        return [], []

    headers = [_text(v) for v in rows[0]]
    if not headers or any(not h for h in headers):
        raise HTTPException(400, f'Sheet "{ws.title}" memiliki header kosong.')
    seen = set()
    duplicates = []
    schema = []
    for h in headers:
        key = _norm(h)
        if key in seen:
            duplicates.append(h)
        else:
            seen.add(key)
            schema.append(h)
    if duplicates:
        raise HTTPException(400, f'Sheet "{ws.title}" memiliki header duplikat: {", ".join(duplicates)}.')

    records = []
    for row in rows[1:]:
        record = {}
        for i, field in enumerate(schema):
            record[field] = _text(row[i]) if i < len(row) else ""
        if any(value for value in record.values()):
            records.append(record)
    return schema, records


async def import_workbook(data: bytes):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(400, f"File Excel tidak valid ({type(exc).__name__}). Gunakan .xlsx.")

    product_ws = None
    for ws in wb.worksheets:
        if _norm(ws.title) == PRODUCT_SHEET_NAME:
            product_ws = ws
            break
    if product_ws is None:
        raise HTTPException(400, 'Sheet wajib bernama "product".')

    rate = await get_rate()
    products, errors = _parse_product_rows(product_ws, rate)
    if not products:
        raise HTTPException(400, "Tidak ada product valid di sheet product.")

    sheet_map = {_norm(ws.title): ws for ws in wb.worksheets if _norm(ws.title) != PRODUCT_SHEET_NAME}
    imported = 0
    updated = 0
    inventory_created = 0
    inventory_skipped = 0
    details = []

    for item in products:
        existing = await db.products.find_one({"name": item["name"]})
        pid = existing["_id"] if existing else str(uuid.uuid4())

        doc = {
            "name": item["name"],
            "description": item["description"],
            "short_description": item["short_description"],
            "price_usd": item["price_usd"],
            "price_idr": item["price_idr"],
            "product_kind": item["product_kind"],
            "delivery_type": "service" if item["product_kind"] == "service" else "inventory",
            "content": "",
            "storage_path": None,
            "original_filename": None,
            "service_wait_minutes": item["wait_minutes"],
            "service_message_template": item["service_message"],
            "active": True,
            "stock": None,
            "stock_mode": "unlimited" if item["product_kind"] == "service" else "auto",
            "manual_stock": None,
            "inventory_enabled": item["product_kind"] == "digital",
            "updated_at": now_iso(),
        }

        if existing:
            await db.products.update_one({"_id": pid}, {"$set": doc})
            updated += 1
        else:
            doc["_id"] = pid
            doc["inventory_schema"] = []
            doc["created_at"] = now_iso()
            await db.products.insert_one(doc)
            imported += 1
            try:
                await notify_product_created(doc)
            except Exception:
                pass

        result = {"name": item["name"], "product_id": pid, "description_generated": True}

        if item["product_kind"] == "digital":
            ws = sheet_map.get(_norm(item["name"]))
            if ws is None:
                errors.append(
                    f'Product "{item["name"]}": sheet inventory dengan nama "{item["name"]}" tidak ditemukan.'
                )
            else:
                schema, records = _inventory_from_sheet(ws)
                if not schema:
                    errors.append(f'Product "{item["name"]}": sheet inventory kosong.')
                else:
                    inv = await add_records(pid, records, schema)
                    inventory_created += inv["created"]
                    inventory_skipped += inv["skipped"]
                    result["schema"] = schema
                    result["inventory_created"] = inv["created"]
                    result["inventory_skipped"] = inv["skipped"]
                    result["stock"] = await available_count(pid)
        else:
            result["stock"] = None

        details.append(result)

    return {
        "ok": True,
        "imported": imported,
        "updated": updated,
        "inventory_created": inventory_created,
        "inventory_skipped": inventory_skipped,
        "errors": errors[:50],
        "details": details,
    }
