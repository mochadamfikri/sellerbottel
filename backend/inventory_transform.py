"""Transform an uploaded inventory spreadsheet without changing stored stock."""
import io
import re
import secrets
import string
from random import SystemRandom

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook

from admin_routes import _parse_inventory_input
from auth import get_current_admin
from db import db

router = APIRouter(prefix="/api/admin/products/{pid}/inventory/transform",
                   dependencies=[Depends(get_current_admin)])

MAX_FILE_BYTES = 5_000_000
MAX_ROWS = 10_000
PASSWORD_SPECIAL = "!@#$%^&*_-"
DOMAIN_RE = re.compile(r"^(?=.{1,253}$)(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}$")


def _is_password_header(header):
    normalized = re.sub(r"[^a-z0-9]", "", header.lower())
    return normalized in {"password", "katasandi", "sandi"}


def _is_email_header(header):
    normalized = header.lower().replace("-", "").replace("_", "").replace(" ", "")
    return "email" in normalized or normalized in {"gmail", "mail"}


def _normalize_domain(value):
    value = str(value or "").strip().lower().lstrip("@")
    if value and not DOMAIN_RE.fullmatch(value):
        raise HTTPException(400, "Domain harus seperti contoh.com, tanpa alamat email lengkap.")
    return value


def random_password():
    groups = (string.ascii_uppercase, string.ascii_lowercase,
              string.digits, PASSWORD_SPECIAL)
    chars = [secrets.choice(group) for group in groups]
    all_chars = "".join(groups)
    chars.extend(secrets.choice(all_chars) for _ in range(8))
    SystemRandom().shuffle(chars)
    return "".join(chars)


async def _read_source(pid: str, upload: UploadFile):
    product = await db.products.find_one({"_id": pid})
    if not product:
        raise HTTPException(404, "Produk tidak ditemukan.")
    if product.get("inventory_mode") == "telegram_session" or product.get("inventory_schema") == ["Session File"]:
        raise HTTPException(400, "File .session tidak dapat diubah sebagai tabel inventory.")
    name = (upload.filename or "").lower()
    if not name.endswith((".xlsx", ".csv", ".txt")):
        raise HTTPException(400, "Gunakan file XLSX, CSV, atau TXT inventory.")
    # The request is already spooled by Starlette; a bounded local read avoids
    # dispatching another threadpool job while parsing the same upload.
    data = upload.file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(400, "File maksimal 5 MB.")
    class InMemoryUpload:
        filename = upload.filename

        async def read(self):
            return data

    clone = InMemoryUpload()
    schema, records = await _parse_inventory_input(clone, "", product)
    if name.endswith(".txt") and records:
        first = records[0]
        if all(str(first.get(field, "")).strip().casefold() == field.casefold() for field in schema):
            records = records[1:]
    if not records:
        raise HTTPException(400, "File tidak berisi baris inventory.")
    if len(records) > MAX_ROWS:
        raise HTTPException(400, "Maksimal 10.000 baris per file.")
    return schema, records


@router.post("/inspect")
async def inspect_inventory_file(pid: str, file: UploadFile = File(...)):
    schema, records = await _read_source(pid, file)
    return {"schema": schema, "row_count": len(records),
            "email_columns": [field for field in schema if _is_email_header(field)],
            "password_columns": [field for field in schema if _is_password_header(field)]}


@router.post("")
async def transform_inventory_file(
    pid: str,
    file: UploadFile = File(...),
    old_domain: str = Form(""),
    new_domain: str = Form(""),
    email_column: str = Form(""),
    password_column: str = Form(""),
    password_mode: str = Form("keep"),
    fixed_password: str = Form(""),
):
    schema, records = await _read_source(pid, file)
    old = _normalize_domain(old_domain)
    new = _normalize_domain(new_domain)
    if bool(email_column) != bool(new):
        raise HTTPException(400, "Pilih kolom email dan isi domain baru untuk mengubah domain.")
    if email_column and (email_column not in schema or not _is_email_header(email_column)):
        raise HTTPException(400, "Kolom email tidak valid.")
    if password_mode not in {"keep", "random", "fixed"}:
        raise HTTPException(400, "Mode kata sandi tidak valid.")
    if password_mode != "keep" and (password_column not in schema or not _is_password_header(password_column)):
        raise HTTPException(400, "File harus punya kolom password, kata sandi, atau sandi.")
    if password_mode == "fixed" and (not fixed_password or len(fixed_password) > 128):
        raise HTTPException(400, "Isi nilai sandi tetap, maksimal 128 karakter.")
    if not new and password_mode == "keep":
        raise HTTPException(400, "Pilih perubahan domain atau kata sandi.")

    wb = Workbook()
    ws = wb.active
    ws.title = "inventory"
    ws.append(schema)
    domain_changes = 0
    password_changes = 0
    for record in records:
        item = dict(record)
        if new:
            email = str(item.get(email_column) or "").strip()
            if email.count("@") == 1 and not any(char.isspace() for char in email):
                local, current = email.rsplit("@", 1)
                if local and current and (not old or current.lower() == old):
                    replacement = f"{local}@{new}"
                    if replacement != email:
                        item[email_column] = replacement
                        domain_changes += 1
        if password_mode != "keep":
            item[password_column] = random_password() if password_mode == "random" else fixed_password
            password_changes += 1
        ws.append([str(item.get(field) or "") for field in schema])
        for cell in ws[ws.max_row]:
            # Preserve uploaded strings that start with '=' as text, not Excel formulas.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(48, max(16, len(str(col[0].value or "")) + 4))
    output = io.BytesIO()
    wb.save(output)
    return Response(output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="inventory-transformed.xlsx"',
                 "Cache-Control": "no-store", "X-Rows": str(len(records)),
                 "X-Domain-Changes": str(domain_changes),
                 "X-Password-Changes": str(password_changes)})
