import io
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from db import db

TZ = ZoneInfo("Asia/Jakarta")
router = APIRouter(prefix="/reports", tags=["admin-reports"])

ORDER_STATUSES = {"pending", "paid", "processing", "delivered", "delivery_failed", "failed", "refunded"}
SALE_STATUSES = {"paid", "processing", "delivered", "delivery_failed", "refunded"}

def _day_start(value):
    return datetime.strptime(value, "%Y-%m-%d").replace(tzinfo=TZ)

def _month_start(value):
    return datetime.strptime(value + "-01", "%Y-%m-%d").replace(tzinfo=TZ)

def _iso(dt):
    return dt.astimezone(timezone.utc).isoformat()

def _period_range(date=None, month=None, start_date=None, end_date=None):
    if start_date:
        start = _day_start(start_date)
    elif month:
        start = _month_start(month)
    elif date:
        start = _day_start(date)
    else:
        now = datetime.now(TZ)
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)

    if end_date:
        end = _day_start(end_date) + timedelta(days=1)
    elif month:
        next_month = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
        end = next_month
    elif date or not start_date:
        end = start + timedelta(days=1)
    else:
        end = start + timedelta(days=1)

    if end <= start:
        raise HTTPException(400, "Rentang tanggal tidak valid.")
    return _iso(start), _iso(end), start, end

def _in_period_query(start_iso, end_iso, field="created_at"):
    return {field: {"$gte": start_iso, "$lt": end_iso}}

async def _load_report(date=None, month=None, start_date=None, end_date=None,
                       currency="all", product_id="", payment_method="all",
                       status="all"):
    start_iso, end_iso, start_dt, end_dt = _period_range(date, month, start_date, end_date)

    order_q = _in_period_query(start_iso, end_iso)
    if currency != "all":
        order_q["currency"] = currency
    if payment_method != "all":
        order_q["payment_method"] = payment_method
    if status != "all":
        order_q["status"] = status
    if product_id:
        order_q["items.product_id"] = product_id

    deposit_q = _in_period_query(start_iso, end_iso)
    if currency != "all":
        deposit_q["currency"] = currency

    orders = await db.purchases.find(order_q).sort("created_at", -1).to_list(50000)
    deposits = await db.deposits.find(deposit_q).sort("created_at", -1).to_list(50000)

    # Only approved deposits count as completed deposit volume.
    approved_deposits = [d for d in deposits if d.get("status") == "approved"]
    refunds = [o for o in orders if o.get("status") == "refunded"]
    sale_orders = [o for o in orders if o.get("status") in SALE_STATUSES]

    summary = {
        "start": start_iso,
        "end": end_iso,
        "orders": len(orders),
        "completed_orders": len([o for o in sale_orders if o.get("status") == "delivered"]),
        "refunded_orders": len(refunds),
        "failed_orders": len([o for o in orders if o.get("status") in {"failed", "delivery_failed"}]),
        "gross_sales": defaultdict(float),
        "refund_total": defaultdict(float),
        "net_sales": defaultdict(float),
        "discount_total": defaultdict(float),
        "deposits": defaultdict(float),
        "unique_buyers": len({o.get("user_tid") for o in sale_orders if o.get("user_tid") is not None}),
        "unique_depositors": len({d.get("user_tid") for d in approved_deposits if d.get("user_tid") is not None}),
        "items_sold": 0,
    }

    for o in sale_orders:
        cur = o.get("currency", "")
        total = float(o.get("total") or 0)
        discount = float(o.get("discount_total") or 0)
        summary["gross_sales"][cur] += total
        if o.get("status") == "refunded":
            summary["refund_total"][cur] += total
        summary["discount_total"][cur] += discount
        if o.get("status") != "refunded":
            summary["items_sold"] += sum(int(i.get("qty") or 0) for i in o.get("items", []))

    for cur, value in summary["gross_sales"].items():
        summary["net_sales"][cur] = value - summary["refund_total"].get(cur, 0)

    for d in approved_deposits:
        cur = d.get("currency", "")
        summary["deposits"][cur] += float(d.get("credited_amount") or d.get("amount") or 0)

    products = defaultdict(lambda: {
        "product_id": "", "name": "", "qty": 0, "gross": defaultdict(float),
        "discount": defaultdict(float), "orders": 0
    })
    for o in sale_orders:
        if o.get("status") == "refunded":
            continue
        seen = set()
        for item in o.get("items", []):
            pid = item.get("product_id") or ""
            if product_id and pid != product_id:
                continue
            key = pid or item.get("name") or "unknown"
            p = products[key]
            p["product_id"] = pid
            p["name"] = item.get("name") or "Unknown"
            qty = int(item.get("qty") or 0)
            p["qty"] += qty
            p["gross"][o.get("currency", "")] += float(item.get("subtotal") or 0)
            p["discount"][o.get("currency", "")] += float(item.get("discount_total") or 0)
            if o.get("_id") not in seen:
                p["orders"] += 1
                seen.add(o.get("_id"))

    daily = defaultdict(lambda: {
        "orders": 0, "gross": defaultdict(float), "deposits": defaultdict(float),
        "refunds": defaultdict(float), "items": 0
    })
    for o in sale_orders:
        key = (o.get("created_at") or "")[:10]
        if not key:
            continue
        daily[key]["gross"][o.get("currency", "")] += float(o.get("total") or 0)
        if o.get("status") == "refunded":
            daily[key]["refunds"][o.get("currency", "")] += float(o.get("total") or 0)
        else:
            daily[key]["orders"] += 1
            daily[key]["items"] += sum(int(i.get("qty") or 0) for i in o.get("items", []))
    for d in approved_deposits:
        key = (d.get("created_at") or "")[:10]
        if key:
            daily[key]["deposits"][d.get("currency", "")] += float(d.get("credited_amount") or d.get("amount") or 0)

    return {
        "summary": dict(summary),
        "orders": orders,
        "deposits": approved_deposits,
        "refunds": refunds,
        "products": dict(products),
        "daily": dict(daily),
        "start_dt": start_dt,
        "end_dt": end_dt,
    }

def _json_summary(summary):
    return {
        **{k: v for k, v in summary.items() if not isinstance(v, defaultdict)},
        "gross_sales": dict(summary["gross_sales"]),
        "refund_total": dict(summary["refund_total"]),
        "net_sales": dict(summary["net_sales"]),
        "discount_total": dict(summary["discount_total"]),
        "deposits": dict(summary["deposits"]),
    }

@router.get("")
async def report_summary(
    date: str = Query("", description="YYYY-MM-DD"),
    month: str = Query("", description="YYYY-MM"),
    start_date: str = "",
    end_date: str = "",
    currency: str = "all",
    product_id: str = "",
    payment_method: str = "all",
    status: str = "all",
):
    if currency not in {"all", "IDR", "USD"}:
        raise HTTPException(400, "Currency tidak valid.")
    if status not in ORDER_STATUSES | {"all"}:
        raise HTTPException(400, "Status tidak valid.")
    data = await _load_report(
        date=date or None, month=month or None, start_date=start_date or None,
        end_date=end_date or None, currency=currency, product_id=product_id,
        payment_method=payment_method, status=status,
    )
    return {
        "range": {"start": data["start_dt"].date().isoformat(), "end": (data["end_dt"] - timedelta(days=1)).date().isoformat()},
        "summary": _json_summary(data["summary"]),
        "products": [
            {"product_id": p["product_id"], "name": p["name"], "qty": p["qty"],
             "orders": p["orders"], "gross": dict(p["gross"]), "discount": dict(p["discount"])}
            for p in data["products"].values()
        ],
        "daily": [
            {"date": day, "orders": row["orders"], "items": row["items"],
             "gross": dict(row["gross"]), "deposits": dict(row["deposits"]), "refunds": dict(row["refunds"])}
            for day, row in sorted(data["daily"].items())
        ],
    }

def _money(value):
    return round(float(value or 0), 2)

def _add_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = __import__("openpyxl").styles.Font(bold=True)
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 100) + 1)]
        width = min(45, max(12, max((len(v) for v in values), default=12) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws

@router.get("/export")
async def export_report(
    date: str = Query(""),
    month: str = Query(""),
    start_date: str = "",
    end_date: str = "",
    currency: str = "all",
    product_id: str = "",
    payment_method: str = "all",
    status: str = "all",
):
    data = await _load_report(
        date=date or None, month=month or None, start_date=start_date or None,
        end_date=end_date or None, currency=currency, product_id=product_id,
        payment_method=payment_method, status=status,
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = data["summary"]
    ws.append(["Metric", "Currency", "Value"])
    for cur in sorted(set(s["gross_sales"]) | set(s["deposits"]) | set(s["refund_total"]) | set(s["net_sales"])):
        ws.append(["Gross Sales", cur, _money(s["gross_sales"].get(cur))])
        ws.append(["Refunds", cur, _money(s["refund_total"].get(cur))])
        ws.append(["Net Sales", cur, _money(s["net_sales"].get(cur))])
        ws.append(["Deposits Approved", cur, _money(s["deposits"].get(cur))])
        ws.append(["Discounts", cur, _money(s["discount_total"].get(cur))])
    ws.append(["Orders", "", s["orders"]])
    ws.append(["Completed Orders", "", s["completed_orders"]])
    ws.append(["Refunded Orders", "", s["refunded_orders"]])
    ws.append(["Failed Orders", "", s["failed_orders"]])
    ws.append(["Unique Buyers", "", s["unique_buyers"]])
    ws.append(["Unique Depositors", "", s["unique_depositors"]])
    ws.append(["Items Sold", "", s["items_sold"]])

    order_rows = []
    for o in data["orders"]:
        products = " | ".join(f'{i.get("name","")} x{i.get("qty",1)}' for i in o.get("items", []))
        order_rows.append([
            o.get("invoice_id"), o.get("_id"), o.get("created_at"), o.get("user_tid"),
            o.get("username"), products, o.get("currency"), o.get("payment_method"),
            o.get("status"), _money(o.get("total")), _money(o.get("discount_total")),
            o.get("coupon_code") or "", o.get("paid_at") or "", o.get("delivered_at") or "",
            o.get("delivery_error") or "",
        ])
    _add_sheet(wb, "Orders",
               ["Invoice", "Order ID", "Created At", "Telegram ID", "Username", "Products",
                "Currency", "Payment Method", "Status", "Total", "Discount", "Coupon", "Paid At", "Delivered At", "Error"],
               order_rows)

    product_rows = []
    for p in data["products"].values():
        for cur, gross in p["gross"].items():
            product_rows.append([p["product_id"], p["name"], cur, p["qty"], p["orders"], _money(gross), _money(p["discount"].get(cur))])
    _add_sheet(wb, "Products", ["Product ID", "Product", "Currency", "Qty Sold", "Orders", "Gross", "Discount"], product_rows)

    deposit_rows = []
    for d in data["deposits"]:
        deposit_rows.append([
            d.get("_id"), d.get("created_at"), d.get("user_tid"), d.get("username"),
            d.get("method"), d.get("coin") or "", d.get("network") or "",
            d.get("currency"), _money(d.get("amount")), _money(d.get("credited_amount") or d.get("amount")),
            d.get("payment_amount") or "", d.get("status"), d.get("tx_hash") or d.get("gopay_tx_id") or "",
            d.get("decided_at") or "",
        ])
    _add_sheet(wb, "Deposits",
               ["Deposit ID", "Created At", "Telegram ID", "Username", "Method", "Coin", "Network",
                "Currency", "Requested", "Credited", "Payment Amount", "Status", "TX", "Decided At"],
               deposit_rows)

    refund_rows = []
    for o in data["refunds"]:
        refund_rows.append([
            o.get("invoice_id"), o.get("_id"), o.get("user_tid"), o.get("username"),
            o.get("currency"), _money(o.get("total")), o.get("refunded_at") or "",
            o.get("refund_reason") or "",
        ])
    _add_sheet(wb, "Refunds", ["Invoice", "Order ID", "Telegram ID", "Username", "Currency", "Amount", "Refunded At", "Reason"], refund_rows)

    daily_rows = []
    for day, row in sorted(data["daily"].items()):
        for cur in sorted(set(row["gross"]) | set(row["deposits"]) | set(row["refunds"])):
            daily_rows.append([day, cur, row["orders"], row["items"], _money(row["gross"].get(cur)),
                               _money(row["deposits"].get(cur)), _money(row["refunds"].get(cur)),
                               _money(row["gross"].get(cur) - row["refunds"].get(cur))])
    _add_sheet(wb, "Daily Breakdown", ["Date", "Currency", "Orders", "Items", "Gross Sales", "Deposits", "Refunds", "Net Sales"], daily_rows)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"report-{data['start_dt'].date().isoformat()}-{(data['end_dt'] - timedelta(days=1)).date().isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
